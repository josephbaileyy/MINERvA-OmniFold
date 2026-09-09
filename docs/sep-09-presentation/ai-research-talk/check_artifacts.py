#!/usr/bin/env python3
"""Check presentation sources, numerical transcription, notes and rendered bounds.

These checks validate the presentation package, not the underlying experiment.
PowerPoint is inspected structurally; PDF is the rendering used for visual review.
"""

from __future__ import annotations

import hashlib
import json
import re
import subprocess
from collections import Counter
from itertools import combinations
from pathlib import Path
from typing import Any
from zipfile import ZipFile

import pymupdf
from pptx import Presentation

HERE = Path(__file__).resolve().parent
REPO = HERE.parents[2]


def require(condition: bool, message: str) -> None:
    """Raise a descriptive error when a check does not hold."""
    if not condition:
        raise ValueError(message)


def check_sources(metrics: dict[str, Any]) -> None:
    """Check receipt hashes and derive the plotted arithmetic from their operands."""
    receipts = []
    for source in metrics["sources"]:
        content = subprocess.check_output(
            ["git", "show", f"{source['revision']}:{source['path']}"], cwd=REPO
        )
        require(
            hashlib.sha256(content).hexdigest() == source["sha256"],
            f"Receipt digest mismatch: {source['path']}",
        )
        receipts.append(json.loads(content))
    agreement, original, repaired = receipts
    science = metrics["scientific"]
    totals = agreement["headline_totals_oi130"]["values"]
    ratio = totals["integral_ours_reported"] / totals["integral_paper_reported"]
    require(abs(ratio - science["integral_ratio"]) < 1e-14, "Integral ratio changed")
    require(
        f"{ratio:.4f}" == "1.0113", "Normalization slide integral ratio needs updating"
    )
    require(
        f"{science['bare_sum_ratio']:.4f}" == "1.0115",
        "Normalization slide bare-sum ratio needs updating",
    )
    require(
        science["agreement_counts"] == [159, 193, 202],
        "Script agreement counts changed",
    )
    require(
        original["gate_B"]["n_pass_gen"] == 1999928, "Checkpoint population changed"
    )
    require(
        repaired["batch_size"] == 512 and repaired["gate_B"]["Bi_max_rel_dev"] == 0,
        "Repaired checkpoint is not the exact matched-batch comparison",
    )
    dates = subprocess.check_output(
        ["git", "log", metrics["revision"], "--no-merges", "--format=%as"],
        cwd=REPO,
        text=True,
    ).splitlines()
    counts = Counter(day[:7] for day in dates)
    for row in metrics["monthly"]:
        require(
            counts[row["month"]] == row["nonmerge_commits"],
            f"Independent monthly census disagrees: {row['month']}",
        )
    require(
        counts["2026-07"] == 263 and counts["2026-08"] == 1970,
        "Script's July/August counts need updating",
    )


def check_additions() -> None:
    """Check the corner plot against its original log and committed evidence."""
    record = json.loads((HERE / "measurements/corner_comparison.json").read_text())
    for source in record["sources"]:
        content = subprocess.check_output(
            ["git", "show", f"{source['revision']}:{source['path']}"], cwd=REPO
        )
        require(
            hashlib.sha256(content).hexdigest() == source["sha256"],
            "Addition source digest changed",
        )
    content = (HERE / record["log_path"]).read_bytes()
    require(
        hashlib.sha256(content).hexdigest() == record["log_sha256"],
        "Corner log digest changed",
    )
    parsed = re.findall(
        r"\[band\] ([^:]+):.*hiE-hiW corner=([\de.+-]+) \(data ([\de.+-]+), data/gen=([\d.]+)\)",
        content.decode(),
    )
    require(
        len(parsed) == len(record["records"]) == 4, "Four-generator population changed"
    )
    for (name, generator, data, ratio), row in zip(
        parsed, record["records"], strict=True
    ):
        require(
            name == row["generator"] and float(ratio) == row["ratio"],
            "Corner point differs from original log",
        )
        require(
            float(generator) == row["generator_integral"]
            and float(data) == row["data_integral"],
            "Corner operands changed",
        )
        require(
            abs(float(data) / float(generator) - float(ratio)) < 0.0006,
            "Rounded corner ratio fails arithmetic check",
        )


def check_layout() -> None:
    """Reject intersecting text boxes and text extending beyond PDF page bounds."""
    elements = json.loads((HERE / "layout.json").read_text())
    for first, second in combinations(elements, 2):
        if first["slide"] != second["slide"]:
            continue
        overlap_x = min(
            first["x"] + first["width"], second["x"] + second["width"]
        ) - max(first["x"], second["x"])
        overlap_y = min(
            first["y"] + first["height"], second["y"] + second["height"]
        ) - max(first["y"], second["y"])
        require(
            not (overlap_x > 1 and overlap_y > 1),
            f"Slide {first['slide']} text overlap: {first['text']!r} / {second['text']!r}",
        )
    document = pymupdf.open(HERE / "research-with-ai.pdf")  # type: ignore[no-untyped-call]
    require(len(document) == 23, "Presentation PDF is not 23 pages")
    for index in range(len(document)):
        page: Any = document[index]
        text = page.get_text()
        require(len(text) > 100, f"PDF page {index + 1} lacks expected text")
        require("\ufffd" not in text, f"Replacement glyph on PDF page {index + 1}")
        for block in page.get_text("dict")["blocks"]:
            if block["type"] != 0:
                continue
            left, top, right, bottom = block["bbox"]
            require(
                left >= 0 and top >= 0 and right <= 960.5 and bottom <= 540.5,
                f"PDF page {index + 1} contains text outside the page",
            )
    script_pdf = pymupdf.open(HERE / "speaker-script.pdf")  # type: ignore[no-untyped-call]
    require(len(script_pdf) == 23, "Printable script does not have one page per slide")


def check_powerpoint() -> None:
    """Check native text, speaker-note pairing and exact embedded plot bytes."""
    presentation = Presentation(str(HERE / "research-with-ai.pptx"))
    sections = re.split(r"(?m)^## ", (HERE / "SCRIPT.md").read_text())[1:]
    require(
        len(presentation.slides) == len(sections) == 23, "Slide / notes count mismatch"
    )
    for index, slide in enumerate(presentation.slides):
        _, notes = sections[index].strip().split("\n", 1)
        require(
            slide.notes_slide.notes_text_frame.text.strip() == notes.strip(),
            f"Notes mismatch on slide {index + 1}",
        )
        require(
            sum(shape.has_text_frame for shape in slide.shapes) >= 5,
            f"Slide {index + 1} is missing editable text",
        )
    with ZipFile(HERE / "research-with-ai.pptx") as archive:
        images = {
            hashlib.sha256(archive.read(name)).hexdigest()
            for name in archive.namelist()
            if name.startswith("ppt/media/")
        }
    for path in (HERE / "figures").glob("*.png"):
        require(
            hashlib.sha256(path.read_bytes()).hexdigest() in images,
            f"PowerPoint lacks the current plot: {path.name}",
        )


def main() -> None:
    """Validate the complete package and write an output digest manifest."""
    metrics = json.loads((HERE / "measurements/metrics.json").read_text())
    check_sources(metrics)
    check_additions()
    from plot_2d_comparison import project, project_ours_errors, read_comparison

    comparison = read_comparison()
    project(comparison, 0)
    project(comparison, 1)
    project_ours_errors(comparison, 0)
    project_ours_errors(comparison, 1)
    from plot_uncertainty_sources import measure_sources

    require(
        measure_sources()
        == json.loads((HERE / "measurements/uncertainty_sources.json").read_text()),
        "Uncertainty source export differs from its input matrices",
    )
    check_layout()
    check_powerpoint()
    files = [
        "research-with-ai.pptx",
        "research-with-ai.pdf",
        "speaker-script.pdf",
        "SCRIPT.md",
    ]
    manifest = {
        name: hashlib.sha256((HERE / name).read_bytes()).hexdigest() for name in files
    }
    (HERE / "artifact-hashes.json").write_text(json.dumps(manifest, indent=2) + "\n")
    print(
        "PASS: source digests, arithmetic, monthly census, 23 slides, paired notes, "
        "editable text, embedded figures, text bounds and printable script"
    )
    print(
        "PDF previews require visual review; native PowerPoint rendering is not exercised here."
    )


if __name__ == "__main__":
    main()
