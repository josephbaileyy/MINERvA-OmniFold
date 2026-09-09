#!/usr/bin/env python3
"""Build an editable PowerPoint and a PDF from one presentation layout.

Speaker notes are read from SCRIPT.md. Scientific plots read the frozen metrics
file; their generation is separate from slide layout. All coordinates are points.
"""

from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any
from xml.sax.saxutils import escape

import pymupdf
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.util import Pt
from reportlab.lib.colors import HexColor
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

HERE = Path(__file__).resolve().parent
WIDTH, HEIGHT = 960, 540
INK = "182A35"
TEAL = "147D83"
ORANGE = "CB653C"
MUTED = "60717A"
PAPER = "FAF9F5"
PALE = "EAF1EF"
WARM = "F5EAE2"
LINE = "D5DEDD"
WHITE = "FFFFFF"


def read_notes() -> list[tuple[str, str]]:
    """Read ordered slide headings and spoken notes from the script."""
    sections = re.split(r"(?m)^## ", (HERE / "SCRIPT.md").read_text())[1:]
    notes = []
    for section in sections:
        title, spoken = section.strip().split("\n", 1)
        notes.append((title, spoken))
    return notes


def register_fonts() -> None:
    """Use the same Arial font for PDF width measurement and slide text."""
    font_root = Path("/System/Library/Fonts/Supplemental")
    if (font_root / "Arial.ttf").exists():
        files = (font_root / "Arial.ttf", font_root / "Arial Bold.ttf")
    else:
        from matplotlib import font_manager

        files = (
            Path(font_manager.findfont("DejaVu Sans")),
            Path(
                font_manager.findfont(
                    font_manager.FontProperties(family="DejaVu Sans", weight="bold")
                )
            ),
        )
    for label, path in zip(("Talk", "Talk-Bold"), files, strict=True):
        pdfmetrics.registerFont(TTFont(label, str(path)))


class Deck:
    """Keep native slide objects and PDF drawings aligned on one canvas.

    Parameters
    ----------
    notes : list of tuple
        Ordered script headings and speaker notes, one entry per slide.
    """

    def __init__(self, notes: list[tuple[str, str]]) -> None:
        self.notes = notes
        self.presentation = Presentation()
        self.presentation.slide_width = Pt(WIDTH)
        self.presentation.slide_height = Pt(HEIGHT)
        self.presentation.core_properties.title = "Research with AI as an undergraduate"
        self.presentation.core_properties.author = "Joseph Bailey"
        self.pdf = canvas.Canvas(
            str(HERE / "research-with-ai.pdf"), pagesize=(WIDTH, HEIGHT)
        )
        self.pdf.setTitle("Research with AI as an undergraduate — Joseph Bailey")
        self.slide: Any = None
        self.index = 0
        self.elements: list[dict[str, Any]] = []

    def rect(self, x: float, y: float, width: float, height: float, fill: str) -> None:
        """Draw an editable rectangle in both formats."""
        shape = self.slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.RECTANGLE, Pt(x), Pt(y), Pt(width), Pt(height)
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(fill)  # type: ignore[no-untyped-call]
        shape.line.fill.background()
        self.pdf.setFillColor(HexColor("#" + fill))
        self.pdf.rect(x, HEIGHT - y - height, width, height, fill=1, stroke=0)

    def text(
        self,
        content: str,
        x: float,
        y: float,
        width: float,
        *,
        size: float = 22,
        color: str = INK,
        bold: bool = False,
        leading: float = 1.22,
    ) -> float:
        """Wrap once, then draw the same editable text lines in both formats."""
        font = "Talk-Bold" if bold else "Talk"
        missing = {
            character
            for character in content
            if not character.isspace()
            and ord(character) not in pdfmetrics.getFont(font).face.charWidths
        }
        if missing:
            raise ValueError(f"Font lacks characters: {missing}")
        lines = []
        for paragraph in content.split("\n"):
            current = ""
            for word in paragraph.split():
                candidate = f"{current} {word}".strip()
                if pdfmetrics.stringWidth(candidate, font, size) > width and current:
                    lines.append(current)
                    current = word
                else:
                    current = candidate
            lines.append(current)
        line_height = size * leading
        height = len(lines) * line_height
        if x < 0 or y < 0 or x + width > WIDTH + 0.1 or y + height > HEIGHT - 5:
            raise ValueError(f"Slide {self.index}: text outside canvas: {content!r}")
        for index, line in enumerate(lines):
            if pdfmetrics.stringWidth(line, font, size) > width + 0.1:
                raise ValueError(f"Unbreakable line too wide: {line!r}")
            top = y + index * line_height
            box = self.slide.shapes.add_textbox(
                Pt(x), Pt(top), Pt(width + 3), Pt(line_height + 3)
            )
            frame = box.text_frame
            frame.margin_top = frame.margin_bottom = 0
            frame.margin_left = frame.margin_right = 0
            frame.word_wrap = False
            paragraph = frame.paragraphs[0]
            paragraph.text = line
            paragraph.font.name = "Arial"
            paragraph.font.size = Pt(size)
            paragraph.font.bold = bold
            paragraph.font.color.rgb = RGBColor.from_string(color)  # type: ignore[no-untyped-call]
            self.pdf.setFont(font, size)
            self.pdf.setFillColor(HexColor("#" + color))
            self.pdf.drawString(x, HEIGHT - top - size * 0.91, line)
        self.elements.append(
            {
                "slide": self.index,
                "text": content,
                "x": x,
                "y": y,
                "width": width,
                "height": height,
                "size": size,
            }
        )
        return height

    def picture(
        self, name: str, x: float, y: float, width: float, height: float
    ) -> None:
        """Fit a figure without changing its aspect ratio."""
        path = HERE / "figures" / f"{name}.png"
        with Image.open(path) as picture:
            factor = min(width / picture.width, height / picture.height)
            drawn_width, drawn_height = picture.width * factor, picture.height * factor
        left, top = x + (width - drawn_width) / 2, y + (height - drawn_height) / 2
        self.slide.shapes.add_picture(
            str(path), Pt(left), Pt(top), Pt(drawn_width), Pt(drawn_height)
        )
        self.pdf.drawImage(
            str(path), left, HEIGHT - top - drawn_height, drawn_width, drawn_height
        )

    def start(self, title: str, section: str, source: str) -> None:
        """Start a slide and attach its matching speaker notes."""
        if self.index:
            self.pdf.showPage()
        self.index += 1
        self.slide = self.presentation.slides.add_slide(
            self.presentation.slide_layouts[6]
        )
        self.rect(0, 0, WIDTH, HEIGHT, PAPER)
        self.rect(0, 0, 12, HEIGHT, TEAL)
        self.text(section.upper(), 44, 24, 860, size=11, color=TEAL, bold=True)
        self.text(title, 44, 52, 872, size=32, bold=True, leading=1.1)
        self.rect(44, 492, 872, 1, LINE)
        self.text(source, 44, 501, 803, size=10, color=MUTED)
        self.text(f"{self.index:02d}", 875, 500, 36, size=12, color=MUTED)
        self.slide.notes_slide.notes_text_frame.text = self.notes[self.index - 1][
            1
        ].strip()

    def card(
        self,
        x: float,
        y: float,
        width: float,
        height: float,
        label: str,
        body: str,
        *,
        fill: str = PALE,
        accent: str = TEAL,
    ) -> None:
        """Draw a short card with a heading and body."""
        self.rect(x, y, width, height, fill)
        label_height = self.text(
            label, x + 18, y + 16, width - 36, size=20, color=accent, bold=True
        )
        body_top = y + 30 + label_height
        body_height = self.text(body, x + 18, body_top, width - 36, size=20)
        if body_top + body_height > y + height - 8:
            raise ValueError(f"Card text exceeds its background: {label}")

    def finish(self) -> None:
        """Write the presentation, PDF, layout inventory, and rendered thumbnails."""
        if self.index != len(self.notes):
            raise ValueError(
                f"{self.index} slides but {len(self.notes)} script sections"
            )
        self.pdf.save()
        self.presentation.save(str(HERE / "research-with-ai.pptx"))
        (HERE / "layout.json").write_text(json.dumps(self.elements, indent=2) + "\n")
        write_script_pdf(self.notes)
        render_preview()


def write_script_pdf(notes: list[tuple[str, str]]) -> None:
    """Create a printable script with one slide's notes per page."""
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer

    document = SimpleDocTemplate(
        str(HERE / "speaker-script.pdf"),
        pagesize=(612, 792),
        leftMargin=48,
        rightMargin=48,
        topMargin=48,
        bottomMargin=48,
    )
    heading = ParagraphStyle(
        "Heading",
        fontName="Talk-Bold",
        fontSize=20,
        leading=25,
        textColor=HexColor("#" + TEAL),
        spaceAfter=20,
    )
    body = ParagraphStyle(
        "Body",
        fontName="Talk",
        fontSize=12,
        leading=17,
        textColor=HexColor("#" + INK),
        spaceAfter=12,
    )
    story = []
    for index, (title, spoken) in enumerate(notes):
        if index:
            story.append(PageBreak())
        story.append(Paragraph(escape(title), heading))
        for paragraph in spoken.strip().split("\n\n"):
            story.append(Paragraph(escape(" ".join(paragraph.splitlines())), body))
        story.append(Spacer(1, 6))
    document.build(story)


def render_preview() -> None:
    """Rasterize the PDF for slide inspection and create a contact sheet."""
    preview = HERE / "preview"
    preview.mkdir(exist_ok=True)
    document = pymupdf.open(HERE / "research-with-ai.pdf")  # type: ignore[no-untyped-call]
    sheet = Image.new("RGB", (1200, ((len(document) + 2) // 3) * 250), "#DDE3E1")
    draw = ImageDraw.Draw(sheet)
    for index in range(len(document)):
        page: Any = document[index]
        page.get_pixmap(matrix=pymupdf.Matrix(1.5, 1.5)).save(  # type: ignore[no-untyped-call]
            preview / f"slide-{index + 1:02d}.png"
        )
        with Image.open(preview / f"slide-{index + 1:02d}.png") as frame:
            frame.thumbnail((388, 218))
            left, top = (index % 3) * 400 + 6, (index // 3) * 250 + 7
            sheet.paste(frame, (left, top))
            draw.text((left + 4, top + 224), f"Slide {index + 1:02d}", fill="#182A35")
    sheet.save(HERE / "slide-overview.png")


def build_content(deck: Deck, metrics: dict[str, Any]) -> None:
    """Lay out the research story, workflow cases, and backup slides."""
    science = metrics["scientific"]
    revision = metrics["revision"][:8]
    deck.start(
        "Research with AI as an undergraduate",
        "Joseph Bailey · September 9, 2026",
        "MINERvA–OmniFold · Nachman group · Personal case study",
    )
    deck.text(
        "What I could build.\nWhat I could trust.\nWhat I understand.",
        58,
        151,
        830,
        size=43,
        bold=True,
        leading=1.2,
    )
    deck.text(
        "A neutrino analysis—and an experiment in how to do research.",
        60,
        370,
        800,
        size=23,
        color=MUTED,
    )
    deck.text(
        "Physics undergraduate · CS coterm · Entering fourth year",
        60,
        441,
        800,
        size=16,
        color=TEAL,
    )

    deck.start(
        "The starting point: an open analysis + OmniFold",
        "01 / The research",
        "MINERvA: arXiv:2106.16210 · Milton et al.: unbinned_unfolding · Start date: speaker account",
    )
    deck.text(
        "“I nodded because I thought I understood.”", 52, 116, 860, size=27, bold=True
    )
    deck.card(
        50,
        175,
        268,
        137,
        "Public MINERvA analysis",
        "Tutorial, data and\na published 2D result",
    )
    deck.card(
        346,
        175,
        268,
        137,
        "An unfolding method",
        "Classifiers learn\nsimulation weights",
    )
    deck.card(
        642,
        175,
        268,
        137,
        "Several LLMs",
        "Explain the pieces;\nhelp put them together",
    )
    deck.text("Simulation truth", 55, 350, 190, size=21, bold=True)
    deck.text("→ detector →", 269, 350, 177, size=21, color=MUTED)
    deck.text("Reconstructed simulation", 474, 350, 340, size=21, bold=True)
    deck.text("Transfer weights back", 57, 402, 330, size=21, color=TEAL)
    deck.text("← classifier compares with data", 421, 402, 472, size=21, color=TEAL)
    deck.text(
        "Initial target: inclusive νμ cross sections in transverse and longitudinal muon momentum.",
        56,
        454,
        850,
        size=16,
        color=MUTED,
    )

    deck.start(
        "How closely does the 2D result reproduce the paper?",
        "01 / The research",
        "Frozen ROOT inputs match committed receipt · Reported phase space · Integrated totals differ by about 1.1%",
    )
    deck.picture("two_d_projections", 36, 121, 888, 325)
    deck.text(
        "Teal: ours ± total error; black/gray: paper total error. Ratio bars: our error / paper central.\nEach uses its own bin widths; ratios match bins. Three pT boundaries differ by 0.005 GeV/c.",
        54,
        450,
        854,
        size=15,
        color=MUTED,
    )

    deck.start(
        "The full 2D comparison shows where differences live",
        "01 / The research",
        "Original 224-bin grid · Positive published statistical-diagonal mask · No significance assigned to residuals",
    )
    deck.picture("two_d_ratio_map", 36, 121, 888, 325)
    deck.text(
        "193 / 205 reported bins are within 10%; three differ by more than 20%. Gray = unreported.\nCorresponding-bin comparison; y labels use this work’s edges. Three pT boundaries differ.",
        54,
        450,
        854,
        size=15,
        color=MUTED,
    )

    deck.start(
        "The agreement survives looking inside the projections",
        "01 / The research",
        "Frozen central values + validated 2D uncertainties · Separate errors, no combined uncertainty on the difference",
    )
    deck.picture("two_d_slices", 36, 121, 888, 325)
    deck.text(
        "Four matched pT slices; all 14 in backup. Density units: cm²/(GeV/c)²/nucleon.\nTeal: ours ± total error; black/gray: paper total error. Ratio bars: our error / paper central.",
        54,
        450,
        854,
        size=15,
        color=MUTED,
    )

    deck.start(
        "Flux and muon reconstruction dominate our errors",
        "01 / The research",
        "Validated 2D source covariances · Each component projected with full bin correlations · No higher-D covariance",
    )
    deck.picture("uncertainty_sources", 36, 119, 888, 309)
    deck.text(
        "Black markers show released paper components with matching scope; blanks do not mean zero.\nValues are medians across projected bins, not additive shares. ML means unfolding-model seed variation.\nMuon energy scale is a subset of reconstruction; its direct comparison is in backup.",
        54,
        437,
        854,
        size=13,
        color=MUTED,
    )

    deck.start(
        "The scope expanded beyond the published reference",
        "01 / The research",
        "2D / 3D / N-D status, governing covariance restrictions, and OI-126 · Current scientific scope",
    )
    stages = [
        ("2D", "Muon momentum", "Central result +\nuncertainty validated", TEAL, PALE),
        ("3D", "+ available energy", "Central result +\nclosure validated", TEAL, PALE),
        ("4D / 5D", "+ q3 and W", "Central results +\nclosures validated", TEAL, PALE),
        (
            "Full event",
            "Particle-cloud inputs",
            "PET diagnostics /\nmethod development",
            ORANGE,
            WARM,
        ),
    ]
    for index, (name, variables, status, color, fill) in enumerate(stages):
        left = 50 + index * 219
        deck.rect(left, 152, 203, 230, fill)
        deck.text(name, left + 15, 171, 177, size=28, bold=True, color=color)
        deck.text(variables, left + 15, 231, 177, size=18)
        deck.text(status, left + 15, 296, 177, size=17, color=color, bold=True)
    deck.text(
        "Higher-dimensional uncertainties remain unadopted.",
        54,
        408,
        850,
        size=25,
        color=ORANGE,
        bold=True,
    )
    deck.text(
        "Beyond reproduction, deciding what would count as validation becomes part of the research.",
        54,
        450,
        850,
        size=18,
    )

    deck.start(
        "Beyond reproduction: a feature the generators miss",
        "01 / The research",
        "Aug 11 four-generator log + ledger VL35–39 · Nine integrated cells · Central-value comparison only",
    )
    deck.text(
        "High available energy and high W: E_avail ≥ 0.8 GeV, W ≥ 1.8 GeV",
        54,
        117,
        850,
        size=21,
        bold=True,
    )
    deck.picture("corner_comparison", 54, 156, 850, 260)
    deck.text(
        "Adding MEC does not close this gap: data / GENIE rises from 1.535 to 1.579.",
        54,
        428,
        850,
        size=20,
        bold=True,
    )
    deck.text(
        "The significance awaits an adopted covariance; no uncertainty bars are shown.",
        54,
        465,
        850,
        size=16,
        color=MUTED,
    )

    deck.start(
        "The workflow became a project of its own",
        "02 / The workflow",
        "Historical routes: orchestration README · MIGRATION-HANDOFF · WAKER · Aug 25 campaign-queue commit",
    )
    milestones = [
        (
            "Start",
            "Ask several models",
            "Explain the tutorial;\nassemble a reproduction.",
        ),
        (
            "July 16–18",
            "Coordinator + workers",
            "Bounded tasks, separate\nverification and review.",
        ),
        (
            "July 19",
            "Deterministic waking",
            "Scripts detect events;\nmodels resume for work.",
        ),
        (
            "August 25",
            "Campaign queue",
            "Stage actions and checks;\nmake execution explicit.",
        ),
    ]
    for index, (date, label, detail) in enumerate(milestones):
        top = 140 + index * 74
        deck.text(date, 54, top, 174, size=20, bold=True, color=TEAL)
        deck.text(label, 239, top, 280, size=21, bold=True)
        deck.text(detail, 546, top, 355, size=19)
    deck.text(
        "These are documented milestones, not clean experimental eras.",
        55,
        462,
        850,
        size=17,
        color=MUTED,
    )

    deck.start(
        "Activity grew. Was the work more effective?",
        "02 / The workflow",
        f"git {revision} · Author months · No merges · One touch = one path in one commit · *Partial months",
    )
    deck.picture("activity", 37, 124, 880, 324)
    deck.text(
        "The right panel measures edit locations. It does not measure hours or scientific value.",
        55,
        455,
        850,
        size=17,
        color=MUTED,
    )

    deck.start(
        "A saved model did not reproduce the event weights",
        "03 / What the checks taught me",
        "Historical PET step-2 receipts · Aug 7 · 1,999,928 pass_gen events · Diagnostic result",
    )
    deck.text(
        "Checkpoint = a saved snapshot of the network’s learned parameters.",
        54,
        116,
        854,
        size=21,
        bold=True,
        color=TEAL,
    )
    deck.picture("checkpoint", 28, 174, 570, 268)
    deck.text("Used in the run", 634, 180, 273, size=19, bold=True)
    deck.text("Last epoch\n(final training pass)", 634, 212, 273, size=21)
    deck.text("Saved to disk", 634, 280, 273, size=19, bold=True)
    deck.text("Best validation epoch\n(lowest held-out loss)", 634, 312, 273, size=21)
    deck.text(
        "Aggregate difference\n< 0.0001", 634, 386, 273, size=21, color=TEAL, bold=True
    )
    deck.text(
        "Learned model parameters → event weights → unfolded distribution",
        54,
        458,
        854,
        size=20,
        bold=True,
    )

    deck.start(
        "The friction: auditing before trying anything",
        "03 / What the workflow cost",
        "Joseph’s account · The checkpoint case shows a useful check; it does not quantify the value of all auditing",
    )
    deck.card(
        50,
        142,
        419,
        251,
        "What slowed me down",
        "Agents wanted to audit\neverything before acting.\n\nI often wanted a small\ntrial-and-error experiment.",
    )
    deck.card(
        491,
        142,
        419,
        251,
        "What I would change",
        "For a cheap, reversible test:\nname the question and try it.\n\nBefore expensive runs or claims:\ncheck the evidence and provenance.",
        fill=WARM,
        accent=ORANGE,
    )
    deck.text(
        "The question is when a check changes the next decision.",
        55,
        421,
        850,
        size=24,
        bold=True,
    )
    deck.text(
        "Proposed workflow change; no measured time saving is claimed.",
        55,
        465,
        850,
        size=16,
        color=MUTED,
    )

    deck.start(
        "My active time shifted into orchestration and routing",
        "03 / What the workflow cost",
        "Joseph’s account · Historical WAKER F2/F3 failures · OI-135 watcher question and implementation decision",
    )
    deck.card(
        50,
        142,
        419,
        251,
        "The work I remember doing",
        "Debugging coordination,\naccount routing and wakeups.\n\nThe logs include wrong executable\npaths and Python environments.",
    )
    deck.card(
        491,
        142,
        419,
        251,
        "A simpler tool can help",
        "“Why do we even need an\nLLM for the watcher?”\n\nA defined job outcome can be\nreported by an ordinary script.",
        fill=WARM,
        accent=ORANGE,
    )
    deck.text(
        "That overhead belongs in any account of productivity.",
        55,
        421,
        850,
        size=24,
        bold=True,
    )
    deck.text(
        "The record supports concrete failures; it does not supply a complete time budget.",
        55,
        465,
        850,
        size=16,
        color=MUTED,
    )

    deck.start(
        "Outside expertise corrected the question itself",
        "04 / Becoming a researcher",
        "COLLABORATOR_QUESTIONS.md · Joseph’s clarification recorded Aug 2 · Verbal guidance, not a written endorsement",
    )
    deck.card(
        50,
        142,
        419,
        248,
        "The mistaken premise",
        "A collaborator question said\nthere was no prior MINERvA\n3D unfolding result.\n\nThat premise was wrong.",
    )
    deck.card(
        491,
        142,
        419,
        248,
        "The corrected framing",
        "MINERvA already had a\n3D unfolding publication.\n\nThe distinction here is the\nunbinned simultaneous method.",
        fill=WARM,
        accent=ORANGE,
    )
    deck.text(
        "A working pipeline does not establish the novelty of its research claim.",
        55,
        417,
        850,
        size=24,
        bold=True,
    )
    deck.text(
        "This correction does not establish who introduced the error or how much AI contributed.",
        55,
        474,
        850,
        size=15,
        color=MUTED,
    )

    deck.start(
        "I have a high-level map. Do I understand the details?",
        "04 / Becoming a researcher",
        "Joseph’s account · Possible benefits of abstraction are hypotheses, not measured learning outcomes",
    )
    deck.card(
        50,
        142,
        419,
        251,
        "My honest answer",
        "AI abstracts so much away\nthat I am unsure what I could\nexplain without it.\n\nI can see the overall approach.",
    )
    deck.card(
        491,
        142,
        419,
        251,
        "The tradeoff I want to test",
        "Abstraction might help me avoid\ngetting stuck in a rabbit hole.\n\nIt might also hide a mistake\nI cannot recognize or explain.",
        fill=WARM,
        accent=ORANGE,
    )
    deck.text(
        "A useful check: predict a change before asking the model.",
        55,
        421,
        850,
        size=24,
        bold=True,
    )
    deck.text(
        "Then explain the result and trace the evidence without relying on the model’s summary.",
        55,
        465,
        850,
        size=16,
        color=MUTED,
    )

    deck.start(
        "Real research output; an unresolved net benefit",
        "05 / What I take into the next project",
        "Personal case study · No matched no-AI baseline, model ranking, human-hours estimate or learning score",
    )
    deck.card(
        50,
        142,
        419,
        251,
        "What I would keep",
        "Freely asking basic questions.\nHelp implementing unfamiliar work.\nExternal references and closure\nas tests of research outputs.",
    )
    deck.card(
        491,
        142,
        419,
        251,
        "What I would change",
        "Smaller, bounded experiments.\nSimpler recurring operations.\nTrack my intervention time.\nPractice explaining key decisions.",
        fill=WARM,
        accent=ORANGE,
    )
    deck.text(
        "What should a student demonstrate before delegating this much?",
        55,
        421,
        850,
        size=24,
        bold=True,
    )
    deck.text(
        "Judge the accepted research result, the total effort, and the student’s understanding separately.",
        55,
        465,
        850,
        size=16,
        color=MUTED,
    )

    deck.start(
        "All 14 transverse-momentum slices",
        "Backup / Full comparison",
        "Backup / Full 2D comparison · Same mask and frozen inputs as the main plots",
    )
    deck.picture("two_d_all_slices", 36, 121, 888, 325)
    deck.text(
        "Full slice inventory; differing pT boundaries labeled. Teal: our total error; black/gray: paper total error.\nDensity units: cm²/(GeV/c)²/nucleon; per-panel powers of ten are printed on the axes.",
        54,
        450,
        854,
        size=15,
        color=MUTED,
    )

    deck.start(
        "The 2D reproduction supplied an external reference",
        "Backup / Agreement counts",
        "Committed agreement receipt, Aug 21 · 205 reported bins · Descriptive central-value comparison",
    )
    deck.picture("agreement", 28, 134, 560, 303)
    deck.text("Integrated cross section", 626, 146, 279, size=20, bold=True)
    deck.text(
        f"{science['total_ours'] / 1e-38:.3f}  this work\n"
        f"{science['total_paper'] / 1e-38:.3f}  published",
        626,
        192,
        279,
        size=27,
        bold=True,
    )
    deck.text("× 10^-38 cm² / nucleon", 626, 270, 279, size=17, color=MUTED)
    deck.text("About 1.1% apart", 626, 316, 279, size=27, color=TEAL, bold=True)
    deck.text(
        "Completed checks: closure, completeness, iterations and standalone uncertainty.",
        52,
        446,
        850,
        size=18,
    )

    deck.start(
        "A reassuring ratio can hide the wrong quantity",
        "Backup / Normalization",
        "2D agreement receipt, Aug 21 · Values re-derived from recorded operands · Ratios rounded",
    )
    deck.text(
        "A differential cross section needs bin areas to become a total.",
        55,
        126,
        850,
        size=24,
    )
    deck.picture("normalization", 62, 178, 833, 62)
    deck.rect(54, 250, 853, 151, PALE)
    deck.text("Calculation", 73, 267, 344, size=20, bold=True)
    deck.text("This work", 472, 267, 176, size=20, bold=True)
    deck.text("Ours / paper", 709, 267, 178, size=20, bold=True)
    deck.text("Integrated total", 73, 313, 344, size=21)
    deck.text("3.073 × 10^-38", 472, 313, 218, size=21)
    deck.text("1.0113", 709, 313, 178, size=21, color=TEAL, bold=True)
    deck.text("Bare sum of densities", 73, 354, 344, size=21, color=ORANGE)
    deck.text("3.732 × 10^-37", 472, 354, 218, size=21, color=ORANGE)
    deck.text("1.0115", 709, 354, 178, size=21, color=ORANGE, bold=True)
    deck.text(
        "Only the integrated row has total-cross-section units: cm²/nucleon.",
        58,
        416,
        850,
        size=17,
        color=MUTED,
    )
    deck.text(
        "Check the definition and units—not only agreement with another number.",
        58,
        451,
        850,
        size=21,
        bold=True,
    )

    deck.start(
        "Repository size describes scope, not understanding",
        "Backup / Size and complexity",
        f"git {revision} · Monthly first-parent snapshots · Tracked .py files · Includes tests, drivers and vendored files",
    )
    deck.picture("scope", 51, 129, 854, 315)
    deck.text(
        "File-edit entropy describes concentration of edits. It is not a code-complexity score.",
        55,
        457,
        850,
        size=18,
        color=MUTED,
    )

    deck.start(
        "What this case study can establish",
        "Backup / Comparison limits",
        "Explicit evidence limits · The absence of a counterfactual does not prevent reporting concrete outcomes",
    )
    deck.card(
        50,
        143,
        419,
        251,
        "Available evidence",
        "Completed research components\nHistorical failure / repair pairs\nRepository activity and scope\nThe speaker's own recollections",
    )
    deck.card(
        491,
        143,
        419,
        251,
        "Not measured",
        "A matched no-AI baseline\nHuman hours saved\nA controlled model ranking\nA learning score",
        fill=WARM,
        accent=ORANGE,
    )
    deck.text(
        "Release dates and commit counts alone cannot identify model-specific productivity gains.",
        55,
        436,
        850,
        size=23,
        bold=True,
    )

    deck.start(
        "Compare the sources released with the paper",
        "Backup / Uncertainty components",
        "Paper ROOT: Total, Flux, MuonEnergyScale, StatOnly covariances · Our validated 2D budget",
    )
    deck.picture("uncertainty_sources_matched", 36, 123, 888, 310)
    deck.text(
        "Our energy-scale group is MINOS + MINERvA; it excludes efficiency, resolution and beam-angle terms.\nEach input uses its own bin widths. Similar source sizes do not establish identical covariance constructions.",
        54,
        444,
        854,
        size=14,
        color=MUTED,
    )

    deck.start(
        "Research references and scope",
        "Backup / Sources",
        "Full source paths, hashes, calculation definitions and editorial notes are in EVIDENCE.md",
    )
    references = [
        ("MINERvA inclusive CC measurement", "Ruterbories et al. · arXiv:2106.16210"),
        ("OmniFold", "Andreassen et al. · arXiv:1911.09107"),
        (
            "Unbinned-unfolding software",
            "Milton et al. · github.com/rymilton/unbinned_unfolding",
        ),
        ("OmniLearn / PET", "Mikuni & Nachman · arXiv:2404.16091"),
        ("MINERvA foundation-model work", "Krzmanc et al. · arXiv:2604.12364"),
    ]
    for index, (label, citation) in enumerate(references):
        top = 134 + index * 57
        deck.text(label, 56, top, 850, size=20, bold=True, color=TEAL)
        deck.text(citation, 56, top + 26, 850, size=17)
    deck.text(
        "2D validated · Higher-D covariance unadopted · PET diagnostic / method development",
        56,
        457,
        850,
        size=17,
        color=ORANGE,
    )


def main() -> None:
    """Build both slide formats and their preview images."""
    register_fonts()
    metrics = json.loads((HERE / "measurements/metrics.json").read_text())
    deck = Deck(read_notes())
    build_content(deck, metrics)
    deck.finish()
    print(
        f"Built {deck.index} slides (16 main + 7 backup), with speaker notes, PDF and previews"
    )


if __name__ == "__main__":
    main()
